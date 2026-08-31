"""Pruebas de métricas y archivos finales."""

from __future__ import annotations

import io
import zipfile

from openpyxl import load_workbook

from reports import (
    build_dashboard_metrics,
    filter_rows,
    make_csv_zip,
    make_excel_report,
    make_pdf_summary,
    monthly_expenses,
)


def report_data() -> dict:
    return {
        "OCR_COLA": [
            {"estado": "CONFIRMADO", "gemini_usado": "FALSE", "gemini_input_tokens": "0", "gemini_output_tokens": "0"},
            {"estado": "NECESITA_REVISION", "gemini_usado": "TRUE", "gemini_input_tokens": "100", "gemini_output_tokens": "12"},
        ],
        "PEAJES": [{"job_id": "1", "fecha": "2026-08-01", "monto_total": "10.50", "moneda": "PEN", "placa": "ABC-123"}],
        "BOLETAS": [{"job_id": "2", "fecha": "2026-08-02", "monto_total": "20,00", "moneda": "PEN", "razon_social": "TIENDA"}],
        "FACTURAS": [{"job_id": "3", "fecha": "2026-07-01", "monto_total": "5.00", "moneda": "USD", "razon_social": "ACME"}],
        "LOGS": [],
    }


def test_dashboard_counts_states_money_and_tokens() -> None:
    metrics = build_dashboard_metrics(report_data())
    assert metrics["jobs_total"] == 2
    assert metrics["confirmados"] == 1
    assert metrics["revision"] == 1
    assert str(metrics["total_pen"]) == "30.50"
    assert metrics["gemini_requests"] == 1
    assert metrics["gemini_input_tokens"] + metrics["gemini_output_tokens"] == 112


def test_monthly_expenses_excludes_other_currency() -> None:
    assert monthly_expenses(report_data()) == [{"mes": "2026-08", "total_pen": 30.5}]


def test_filters_search_and_date() -> None:
    rows = report_data()["PEAJES"] + report_data()["BOLETAS"]
    assert [row["job_id"] for row in filter_rows(rows, query="abc", date_from="2026-08-01")] == ["1"]


def test_excel_pdf_and_zip_are_valid_files() -> None:
    data = report_data()
    excel = make_excel_report(data)
    workbook = load_workbook(io.BytesIO(excel), read_only=True)
    assert {"RESUMEN", "PEAJES", "BOLETAS", "FACTURAS", "OCR_COLA", "LOGS"} <= set(workbook.sheetnames)
    pdf = make_pdf_summary(data)
    assert pdf.startswith(b"%PDF-")
    csv_zip = make_csv_zip(data)
    with zipfile.ZipFile(io.BytesIO(csv_zip)) as archive:
        assert "peajes.csv" in archive.namelist()
