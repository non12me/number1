"""Métricas, filtros y exportaciones de la versión final."""

from __future__ import annotations

import csv
import io
import json
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Iterable


RESULT_SHEETS = ("PEAJES", "BOLETAS", "FACTURAS")


def safe_decimal(value: Any) -> Decimal:
    text = str(value or "").strip().replace(" ", "")
    if not text:
        return Decimal("0")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".") if text.rfind(",") > text.rfind(".") else text.replace(",", "")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def build_dashboard_metrics(data: dict[str, list[dict[str, str]]]) -> dict[str, Any]:
    queue = data.get("OCR_COLA", [])
    results = [row for name in RESULT_SHEETS for row in data.get(name, [])]
    states = Counter(row.get("estado", "SIN_ESTADO") for row in queue)
    total_pen = sum(
        (safe_decimal(row.get("monto_total")) for row in results if row.get("moneda", "PEN") in {"", "PEN"}),
        Decimal("0"),
    )
    input_tokens = sum(int(row.get("gemini_input_tokens") or 0) for row in queue)
    output_tokens = sum(int(row.get("gemini_output_tokens") or 0) for row in queue)
    confirmed = states.get("CONFIRMADO", 0)
    return {
        "jobs_total": len(queue),
        "pendientes": states.get("PENDIENTE", 0),
        "procesando": states.get("PROCESANDO", 0),
        "revision": sum(states.get(name, 0) for name in ("EXTRAIDO_LOCAL", "NECESITA_REVISION", "PENDIENTE_RESULTADO")),
        "confirmados": confirmed,
        "rechazados": states.get("RECHAZADO", 0),
        "errores": states.get("ERROR", 0),
        "resultados": len(results),
        "total_pen": total_pen,
        "gemini_input_tokens": input_tokens,
        "gemini_output_tokens": output_tokens,
        "gemini_requests": sum(1 for row in queue if row.get("gemini_usado", "").upper() == "TRUE"),
        "confirmation_rate": confirmed / len(queue) if queue else 0.0,
    }


def type_summary(data: dict[str, list[dict[str, str]]]) -> list[dict[str, Any]]:
    return [
        {
            "tipo": name,
            "documentos": len(data.get(name, [])),
            "total_pen": float(sum(
                (safe_decimal(row.get("monto_total")) for row in data.get(name, []) if row.get("moneda", "PEN") in {"", "PEN"}),
                Decimal("0"),
            )),
        }
        for name in RESULT_SHEETS
    ]


def monthly_expenses(data: dict[str, list[dict[str, str]]]) -> list[dict[str, Any]]:
    totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    for sheet in RESULT_SHEETS:
        for row in data.get(sheet, []):
            if row.get("moneda", "PEN") not in {"", "PEN"}:
                continue
            try:
                month = datetime.strptime(row.get("fecha", ""), "%Y-%m-%d").strftime("%Y-%m")
            except ValueError:
                continue
            totals[month] += safe_decimal(row.get("monto_total"))
    return [{"mes": month, "total_pen": float(totals[month])} for month in sorted(totals)]


def filter_rows(
    rows: Iterable[dict[str, str]],
    query: str = "",
    date_from: str = "",
    date_to: str = "",
) -> list[dict[str, str]]:
    needle = query.strip().casefold()
    result = []
    for row in rows:
        date = row.get("fecha", "")
        if date_from and date and date < date_from:
            continue
        if date_to and date and date > date_to:
            continue
        if needle and needle not in " ".join(str(value) for value in row.values()).casefold():
            continue
        result.append(row)
    return result


def _export_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value


def make_excel_report(data: dict[str, list[dict[str, str]]]) -> bytes:
    """Crea un XLSX autocontenido con resumen y pestañas operativas."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    summary = workbook.active
    summary.title = "RESUMEN"
    metrics = build_dashboard_metrics(data)
    summary.append(["Indicador", "Valor"])
    for key, value in metrics.items():
        summary.append([key, float(value) if isinstance(value, Decimal) else value])

    export_sheets = ("PEAJES", "BOLETAS", "FACTURAS", "OCR_COLA", "LOGS")
    for name in export_sheets:
        worksheet = workbook.create_sheet(name)
        rows = data.get(name, [])
        headers = list(rows[0]) if rows else []
        if headers:
            worksheet.append(headers)
            for row in rows:
                worksheet.append([_export_value(row.get(header, "")) for header in headers])
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.freeze_panes = "A2"
        else:
            worksheet.append(["Sin registros"])

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="155E75")
            cell.alignment = Alignment(horizontal="center")
        for column in worksheet.columns:
            width = min(45, max(12, *(len(str(cell.value or "")) + 2 for cell in column[:100])))
            worksheet.column_dimensions[column[0].column_letter].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_csv_zip(data: dict[str, list[dict[str, str]]]) -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name in (*RESULT_SHEETS, "OCR_COLA", "LOGS"):
            rows = data.get(name, [])
            text = io.StringIO(newline="")
            headers = list(rows[0]) if rows else ["sin_registros"]
            writer = csv.DictWriter(text, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            for row in rows:
                writer.writerow({key: _export_value(value) for key, value in row.items()})
            archive.writestr(f"{name.lower()}.csv", "\ufeff" + text.getvalue())
    return output.getvalue()


def make_pdf_summary(data: dict[str, list[dict[str, str]]]) -> bytes:
    """Crea un PDF breve con KPIs y totales; no incluye datos sensibles línea por línea."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle, Paragraph

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=A4, title="Resumen OCR documental")
    styles = getSampleStyleSheet()
    metrics = build_dashboard_metrics(data)
    rows = [
        ["Indicador", "Valor"],
        ["Jobs totales", str(metrics["jobs_total"])],
        ["Confirmados", str(metrics["confirmados"])],
        ["Pendientes de revisión", str(metrics["revision"])],
        ["Errores", str(metrics["errores"])],
        ["Resultados", str(metrics["resultados"])],
        ["Total PEN", f"S/ {metrics['total_pen']:.2f}"],
        ["Solicitudes Gemini", str(metrics["gemini_requests"])],
        ["Tokens Gemini", str(metrics["gemini_input_tokens"] + metrics["gemini_output_tokens"])],
    ]
    table = Table(rows, colWidths=[240, 180])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155E75")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("PADDING", (0, 0), (-1, -1), 8),
    ]))
    document.build([
        Paragraph("OCR Documental - Resumen", styles["Title"]),
        Paragraph(datetime.now().strftime("Generado el %Y-%m-%d %H:%M"), styles["Normal"]),
        Spacer(1, 18),
        table,
    ])
    return output.getvalue()
